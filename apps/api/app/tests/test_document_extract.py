from io import BytesIO

import fitz
from docx import Document
from fastapi.testclient import TestClient

from app.core.config import settings
from app.main import app
from app.services.documents.extract import extract_document
from app.services.documents.text_utils import looks_like_heading, split_into_sections

client = TestClient(app)
AUTH = {"X-Internal-Token": settings.internal_api_token}


def _pdf_bytes(pages: list[str]) -> bytes:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Chapter 1 Foundations", level=1)
    document.add_paragraph("Leadership begins with character.")
    document.add_heading("Chapter 2 Practice", level=1)
    document.add_paragraph("Practice turns knowledge into wisdom.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_looks_like_heading_detects_chapters() -> None:
    assert looks_like_heading("CHAPTER 1 BEGINNINGS")
    assert looks_like_heading("1. Introduction")
    assert not looks_like_heading("This is a normal sentence about teaching.")


def test_split_into_sections_keeps_page_ranges() -> None:
    sections = split_into_sections(
        [
            ("CHAPTER 1 START\nWelcome to the course.", 1),
            ("More chapter one content.", 2),
            ("CHAPTER 2 NEXT\nSecond chapter body.", 3),
        ]
    )
    assert len(sections) == 2
    assert sections[0]["page_start"] == 1
    assert sections[0]["page_end"] == 2
    assert sections[1]["page_start"] == 3


def test_extract_pdf_returns_sections() -> None:
    pdf = _pdf_bytes(
        [
            "CHAPTER 1 VISION\nCreators need clarity.",
            "CHAPTER 2 CRAFT\nClarity becomes curriculum.",
        ]
    )
    result = extract_document(
        file_bytes=pdf,
        file_type="pdf",
        original_filename="course.pdf",
    )
    assert result["page_count"] == 2
    assert result["section_count"] >= 1
    assert result["sections"][0]["extracted_text"]


def test_extract_docx_and_txt() -> None:
    docx_result = extract_document(
        file_bytes=_docx_bytes(),
        file_type="docx",
        original_filename="notes.docx",
    )
    assert docx_result["section_count"] >= 2

    txt_result = extract_document(
        file_bytes=b"SECTION 1 OPENING\n\nHello teachers.\n\nSECTION 2 CLOSE\n\nGoodbye.",
        file_type="txt",
        original_filename="notes.txt",
    )
    assert txt_result["section_count"] >= 2


def test_extract_endpoint_requires_token() -> None:
    pdf = _pdf_bytes(["Hello from Luminary."])
    denied = client.post(
        "/api/v1/documents/extract",
        files={"file": ("sample.pdf", pdf, "application/pdf")},
        data={"file_type": "pdf", "original_filename": "sample.pdf"},
    )
    assert denied.status_code == 401

    ok = client.post(
        "/api/v1/documents/extract",
        headers=AUTH,
        files={"file": ("sample.pdf", pdf, "application/pdf")},
        data={"file_type": "pdf", "original_filename": "sample.pdf"},
    )
    assert ok.status_code == 200
    body = ok.json()
    assert body["section_count"] >= 1
    assert body["sections"][0]["section_title"]
